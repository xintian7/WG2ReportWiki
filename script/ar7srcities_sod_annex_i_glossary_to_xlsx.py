#!/usr/bin/env python3
"""Extract the SRCities Annex I glossary to an Excel workbook.

Indented glossary terms are subterms of the most recent non-indented term and
are written with that main term in the ``Parent`` column.

Example:
    /opt/anaconda3/envs/tsu/bin/python script/annex_i_glossary_to_xlsx.py
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


REFERENCE_STYLE = "srcities references"
SOURCE_NAME = "SRCities-SOD"


@dataclass
class GlossaryEntry:
    term: str
    explanation: str
    parent: str = ""


def normalize_space(text: str) -> str:
    """Collapse Word's line and non-breaking-space variants into plain spaces."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def is_term(paragraph) -> bool:
    """Return whether a paragraph consists entirely of bold term text."""
    non_empty_runs = [run for run in paragraph.runs if normalize_space(run.text)]
    return bool(non_empty_runs) and all(run.bold is True for run in non_empty_runs)


def is_subterm(paragraph) -> bool:
    """Return whether the term is indented beneath its current main term."""
    indent = paragraph.paragraph_format.left_indent
    return indent is not None and indent.pt > 0


def glossary_paragraphs(input_docx: Path):
    """Yield glossary paragraphs, stopping before the reference list."""
    document = Document(str(input_docx))
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.strip().casefold() == REFERENCE_STYLE:
            break
        yield paragraph


def extract_entries(input_docx: Path) -> list[GlossaryEntry]:
    """Extract de-duplicated glossary entries and their hierarchical parents."""
    entries: list[GlossaryEntry] = []
    current_entry: GlossaryEntry | None = None
    current_parent = ""

    for paragraph in glossary_paragraphs(input_docx):
        text = normalize_space(paragraph.text)
        if not text:
            continue

        if is_term(paragraph):
            if current_entry is not None:
                entries.append(current_entry)

            parent = current_parent if is_subterm(paragraph) else ""
            current_entry = GlossaryEntry(term=text, explanation="", parent=parent)
            if not parent:
                current_parent = text
            continue

        if current_entry is not None:
            current_entry.explanation = normalize_space(f"{current_entry.explanation} {text}")

    if current_entry is not None:
        entries.append(current_entry)

    unique_entries: dict[str, GlossaryEntry] = {}
    for entry in entries:
        key = entry.term.casefold()
        existing_entry = unique_entries.get(key)
        if existing_entry is None or (
            existing_entry.explanation.casefold().startswith("see ")
            and not entry.explanation.casefold().startswith("see ")
        ):
            unique_entries[key] = entry

    return list(unique_entries.values())


def write_workbook(entries: list[GlossaryEntry], output_path: Path) -> None:
    """Write the glossary entries to a formatted Excel workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Glossary"
    worksheet.append(["Term", "Explanation", "Parent", "Source"])

    for entry in entries:
        worksheet.append([entry.term, entry.explanation, entry.parent, SOURCE_NAME])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for column in worksheet.iter_cols(min_row=2):
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.column_dimensions["A"].width = 44
    worksheet.column_dimensions["B"].width = 120
    worksheet.column_dimensions["C"].width = 44
    worksheet.column_dimensions["D"].width = 20
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract the SRCities Annex I glossary to Excel.")
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("data/Glossary/SRCities_SOD_AnnexI_Final.docx"),
        help="Path to the Annex I DOCX glossary.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/Glossary/AR7SOD_Glossary.xlsx"),
        help="Path for the generated Excel workbook.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise FileNotFoundError(f"Input DOCX not found: {args.input}")

    entries = extract_entries(args.input)
    if not entries:
        raise RuntimeError("No glossary entries were extracted from the DOCX.")

    write_workbook(entries, args.output)
    print(f"Wrote {len(entries)} glossary entries to {args.output}")


if __name__ == "__main__":
    main()